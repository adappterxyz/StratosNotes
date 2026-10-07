"""
Builds landing/public/deck.pptx (served at sp.stratoslab.app/deck.pptx), the TOKEN2049 Origins pitch deck.
Every figure comes from docs/SUBMISSION.md and docs/crosschain.md.

    python3 docs/deck/build_deck.py
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parents[2] / "landing" / "public" / "deck.pptx"

# Palette: landing/src/notes.css (dark ledger glass, emerald / teal / violet).
BG = RGBColor(0x05, 0x0B, 0x0A)
BG2 = RGBColor(0x0C, 0x16, 0x13)
CARD = RGBColor(0x10, 0x1C, 0x18)
LINE = RGBColor(0x24, 0x34, 0x2F)
TEXT = RGBColor(0xE9, 0xF3, 0xEF)
SOFT = RGBColor(0xB4, 0xC7, 0xC0)
MUTED = RGBColor(0x7A, 0x92, 0x8A)
EMERALD = RGBColor(0x10, 0xB9, 0x81)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)  # side margin

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------
def slide(notes: str):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, runs, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6, name=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, opts)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        parts = [(p, {})] if isinstance(p, str) else p
        for t, o in parts:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def title(s, t, kicker=None):
    if kicker:
        text(s, M, Inches(0.45), Inches(10), Inches(0.3), kicker.upper(), size=12, color=EMERALD, bold=True)
    text(s, M, Inches(0.75), W - 2 * M, Inches(0.9), t, size=32, bold=True, name="Title")


def card(s, x, y, w, h, fill=CARD, line=LINE, name=None):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.adjustments[0] = 0.08 if min(w, h) > Inches(1.2) else 0.18
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
        r.line.width = Pt(1)
    r.shadow.inherit = False
    if name:
        r.name = name
    r.text_frame.text = ""
    return r


def dot(s, x, y, color, d=Inches(0.16)):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    o.fill.solid()
    o.fill.fore_color.rgb = color
    o.line.fill.background()
    o.shadow.inherit = False
    return o


def arrow(s, x1, y1, x2, y2, color=MUTED, width=1.75, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def box(s, x, y, w, h, head, sub=None, accent=EMERALD, fill=CARD, head_size=15, sub_size=12):
    card(s, x, y, w, h, fill=fill, line=accent)
    pad = Inches(0.14)
    paras = [[(head, {"bold": True, "size": head_size})]]
    if sub:
        paras.append([(sub, {"size": sub_size, "color": SOFT})])
    text(s, x + pad, y, w - 2 * pad, h, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)


def table(s, x, y, w, rows, col_w, size=13, head_color=EMERALD, row_h=Inches(0.42)):
    shp = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
    tbl = shp.table
    # drop the built-in style banding so our fills win
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = cw
    for r, row in enumerate(rows):
        tbl.rows[r].height = row_h
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG2 if r == 0 else (CARD if r % 2 else BG)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (c == len(row) - 1 and r > 0 and val[:1].isdigit()) else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = val
            run.font.name = FONT
            run.font.size = Pt(size - 1 if r == 0 else size)
            run.font.bold = r == 0
            run.font.color.rgb = head_color if r == 0 else TEXT
    return shp


def footer(s, n):
    text(s, M, H - Inches(0.45), Inches(6), Inches(0.25), "StratosNotes · sp.stratoslab.app", size=10, color=MUTED)
    text(s, W - M - Inches(1), H - Inches(0.45), Inches(1), Inches(0.25), str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


# ---------- 1. Title ----------
s = slide(
    "A structured note is a promise about dates and prices. Today a bank's calculation agent watches the "
    "dates and a paying agent moves the money, and you can only buy the note where the bank is. StratosNotes "
    "is a self-service issuance engine and marketplace: the note runs on Solana, Chainlink CRE is the "
    "calculation and paying agent, and investors can be on Solana or Ethereum."
)
text(s, M, Inches(1.0), Inches(6), Inches(0.35), "TOKEN2049 ORIGINS", size=14, color=EMERALD, bold=True)
text(s, M, Inches(1.9), Inches(11), Inches(1.4), "StratosNotes", size=72, bold=True)
text(s, M, Inches(3.35), Inches(9.6), Inches(1.4),
     "Structured notes anyone can issue in minutes. The note runs on Solana, Chainlink CRE observes and pays it, "
     "and investors buy on Solana or Ethereum.", size=24, color=SOFT)
chips = [("Solana", EMERALD), ("Chainlink CRE", TEAL), ("Chainlink CCIP", VIOLET)]
cx = M
for label, col in chips:
    cw = Inches(0.45 + 0.13 * len(label))
    card(s, cx, Inches(5.05), cw, Inches(0.48), fill=BG2, line=col)
    dot(s, cx + Inches(0.17), Inches(5.21), col)
    text(s, cx + Inches(0.42), Inches(5.05), cw - Inches(0.5), Inches(0.48), label, size=15, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    cx += cw + Inches(0.2)
text(s, M, Inches(6.35), Inches(8), Inches(0.35),
     [[("sp.stratoslab.app", {"bold": True, "color": TEXT}), ("   ·   Solana devnet + Ethereum Sepolia", {"color": MUTED})]],
     size=16)

# ---------- 2. Problem ----------
s = slide(
    "Read the example coupon aloud. Someone has to watch those dates, read those prices and move the money. "
    "Today that is a bank's calculation agent and paying agent; creating a note takes weeks and a legal team; "
    "and investors can only buy it where the issuer's platform lives."
)
title(s, "A note is a promise about dates and prices", "The problem")
card(s, M, Inches(2.0), Inches(5.6), Inches(4.4), fill=BG2)
text(s, M + Inches(0.4), Inches(2.35), Inches(4.8), Inches(3.8), [
    [("“", {"size": 60, "color": EMERALD, "bold": True})],
    [("Every quarter, if the worst of ETH and BTC is at or above 70% of its starting level you get a 2.5% "
      "coupon; if both are at or above 100% you get your money back early.", {"size": 20})],
], space_after=0)
items = [
    ("A calculation agent", "watches the dates and prices", EMERALD),
    ("A paying agent", "moves the money, and investors trust both", TEAL),
    ("Weeks and a legal team", "to create one note", VIOLET),
    ("One venue", "investors buy only where the issuer's platform lives", MUTED),
]
y = Inches(2.0)
for head, sub, col in items:
    card(s, Inches(6.75), y, Inches(5.98), Inches(0.98))
    dot(s, Inches(7.05), y + Inches(0.41), col)
    text(s, Inches(7.45), y + Inches(0.14), Inches(5.1), Inches(0.75),
         [[(head, {"bold": True, "size": 18})], [(sub, {"size": 14, "color": SOFT})]], space_after=0)
    y += Inches(1.14)
footer(s, 2)

# ---------- 3. What it does ----------
s = slide(
    "One workspace, in the design of Flow. Start from a published template, a term sheet, a sentence to the AI "
    "or a blank canvas. Save any design to the library. Issue: set size, strike and observation dates, deposit "
    "the coupon reserve, publish. Sell on either chain: Solana investors swap USDC for units in one "
    "transaction; Ethereum investors send tUSD and a subscribe call over CCIP. On every date CRE reads the "
    "Chainlink feeds and writes a signed report; coupons, autocall, knock-in and redemption settle on-chain."
)
title(s, "From idea to paid coupons, in five steps", "What it does")
steps = [
    ("Design", "Template, term sheet, AI or BPMN canvas. Validated as you edit."),
    ("Template", "Save to the library: published on Solana by its hash, signed by the author."),
    ("Issue", "Size, strike and observation dates, coupon reserve. The book opens."),
    ("Sell", "Solana: USDC for units in one transaction. Ethereum: tUSD over CCIP."),
    ("Observe & pay", "CRE reads Chainlink feeds and settles on-chain. Paid on either chain."),
]
n = len(steps)
gap = Inches(0.3)
cw = int((W - 2 * M - gap * (n - 1)) / n)
cols = [EMERALD, EMERALD, TEAL, TEAL, VIOLET]
for i, (head, sub) in enumerate(steps):
    x = M + i * (cw + gap)
    card(s, x, Inches(2.2), cw, Inches(3.5))
    text(s, x + Inches(0.2), Inches(2.4), cw - Inches(0.5), Inches(0.75), str(i + 1), size=36, bold=True, color=cols[i])
    text(s, x + Inches(0.2), Inches(3.3), cw - Inches(0.4), Inches(0.5), head, size=18, bold=True)
    text(s, x + Inches(0.2), Inches(3.9), cw - Inches(0.4), Inches(1.7), sub, size=14, color=SOFT)
    if i < n - 1:
        arrow(s, x + cw + Inches(0.03), Inches(3.95), x + cw + gap - Inches(0.03), Inches(3.95), color=MUTED, width=2)
text(s, M, Inches(5.95), W - 2 * M, Inches(0.65),
     "Five products: fixed coupon note, reverse convertible, phoenix with memory, snowball, principal-protected. "
     "One underlying or a worst-of basket of up to three. Cash or physical settlement.", size=14, color=MUTED)
footer(s, 3)

# ---------- 4. Architecture ----------
s = slide(
    "Left to right. Any authoring route becomes a BPMN workflow, compiled to a definition stored on Solana once, "
    "by its hash. One Anchor program, flow_engine, runs every issuance. Above it, the CRE notes-keeper reads "
    "Chainlink feeds for ETH, BTC and SOL and writes signed reports into the engine through the forwarder. To the "
    "right, CCIP: Ethereum investors subscribe in through ccip_receive; payouts CRE locked in the outbox go out "
    "through a permissionless relay and ccip_send. The relay decides nothing."
)
title(s, "One program, one CRE workflow, one CCIP lane", "How it works")
bw, bh = Inches(2.15), Inches(0.95)
# authoring column
ax = M
box(s, ax, Inches(2.0), bw, bh, "Author", "template · term sheet · AI · canvas", accent=LINE)
box(s, ax, Inches(3.55), bw, bh, "BPMN workflow", "validated, compiled", accent=LINE)
box(s, ax, Inches(5.1), bw, bh, "Definition", "Solana account, by hash", accent=LINE)
arrow(s, ax + bw / 2, Inches(2.95), ax + bw / 2, Inches(3.55))
arrow(s, ax + bw / 2, Inches(4.5), ax + bw / 2, Inches(5.1))
# engine
ex, ey, ew, eh = Inches(3.45), Inches(3.35), Inches(3.3), Inches(2.15)
box(s, ex, ey, ew, eh, "flow_engine", "One Anchor program on Solana.\nRuns every issuance: state machine,\nholdings, vault, outbox",
    accent=EMERALD, fill=BG2, head_size=20, sub_size=13)
arrow(s, ax + bw, Inches(5.57), ex, Inches(5.0))
# CRE + feeds
box(s, Inches(3.45), Inches(1.85), Inches(1.55), Inches(0.95), "Chainlink feeds", "ETH · BTC · SOL", accent=TEAL, head_size=14)
box(s, Inches(5.25), Inches(1.85), Inches(1.5), Inches(0.95), "CRE keeper", "observe · decide", accent=TEAL, head_size=14)
arrow(s, Inches(5.0), Inches(2.32), Inches(5.25), Inches(2.32), color=TEAL)
arrow(s, Inches(6.0), Inches(2.8), Inches(6.0), ey, color=TEAL)
text(s, Inches(6.08), Inches(2.88), Inches(1.5), Inches(0.4), "signed reports", size=11, color=TEAL)
# outbound
ox = Inches(7.35)
box(s, ox, Inches(3.35), Inches(1.45), Inches(0.9), "Outbox", "locked by CRE", accent=VIOLET, head_size=14)
box(s, Inches(9.15), Inches(3.35), Inches(1.45), Inches(0.9), "Relay", "anyone can run", accent=VIOLET, head_size=14)
arrow(s, ex + ew, Inches(3.8), ox, Inches(3.8), color=VIOLET)
arrow(s, Inches(8.8), Inches(3.8), Inches(9.15), Inches(3.8), color=VIOLET)
# CCIP + Ethereum
cx_ = Inches(10.95)
box(s, cx_, Inches(3.35), Inches(1.78), Inches(2.15), "Chainlink CCIP", "tUSD · tETH\ntBTC · tSOL", accent=VIOLET, head_size=15)
arrow(s, Inches(10.6), Inches(3.8), cx_, Inches(3.8), color=VIOLET)
box(s, cx_, Inches(1.85), Inches(1.78), Inches(0.95), "Ethereum", "investors (Sepolia)", accent=VIOLET, head_size=15)
arrow(s, Inches(11.6), Inches(3.35), Inches(11.6), Inches(2.8), color=VIOLET)
arrow(s, Inches(12.1), Inches(2.8), Inches(12.1), Inches(3.35), color=VIOLET, dash=True)
# inbound back into the engine
arrow(s, cx_, Inches(5.05), ex + ew, Inches(5.05), color=VIOLET, dash=True)
text(s, Inches(7.35), Inches(5.12), Inches(3.5), Inches(0.35), "subscriptions in: ccip_receive", size=11, color=VIOLET)
text(s, Inches(7.35), Inches(4.3), Inches(3.5), Inches(0.35), "payouts out: ccip_send", size=11, color=VIOLET)
text(s, M, Inches(6.45), W - 2 * M, Inches(0.4),
     "Prices always come from Chainlink feeds; the lifecycle always runs on Solana; money and delivered tokens cross chains.",
     size=13, color=MUTED)
footer(s, 4)

# ---------- 5. Solana ----------
s = slide(
    "Best Use of Solana. One Anchor program runs BPMN workflows on-chain. A compiled workflow is stored once, "
    "addressed by its hash, and every issuance is a process of it. Five product engines, worst-of baskets, "
    "physical delivery, the template library and any workflow edited on the canvas all run on the same program. "
    "It has its own heap allocator so large definitions fit Solana's 32 KiB heap on the report path."
)
title(s, "Every note is a BPMN workflow on one Solana program", "Best Use of Solana")
feats = [
    ("State machine", "Token multiset; gateways on price predicates (min / max for worst-of); fork / join"),
    ("Time and roles", "Timers on per-issuance dates, open roles, repeatable subscription windows"),
    ("Holdings and DvP", "Internal holdings ledger, atomic USDC-for-note swap, transfers in whole or in part"),
    ("Distribution", "Pays every holder in cash or the underlying's token, from a token vault per issuance"),
    ("CCIP native", "A CCIP receiver and sender: subscribe in with ccip_receive, pay out with ccip_send"),
    ("Fits the runtime", "Own heap allocator so large definitions fit Solana's 32 KiB heap"),
]
gw, gh = Inches(3.85), Inches(1.75)
gx, gy = Inches(0.3), Inches(0.3)
for i, (head, sub) in enumerate(feats):
    x = M + (i % 3) * (gw + gx)
    y = Inches(2.0) + (i // 3) * (gh + gy)
    card(s, x, y, gw, gh)
    dot(s, x + Inches(0.3), y + Inches(0.36), EMERALD)
    text(s, x + Inches(0.6), y + Inches(0.25), gw - Inches(0.85), Inches(0.4), head, size=18, bold=True)
    text(s, x + Inches(0.3), y + Inches(0.72), gw - Inches(0.6), gh - Inches(0.8), sub, size=14, color=SOFT)
text(s, M, Inches(6.15), W - 2 * M, Inches(0.4),
     [[("flow_engine", {"bold": True, "color": TEXT}),
       ("  ·  Anchor 0.31.1  ·  devnet 9a5xpgRgK7NQMVtYvLuVq1XooK3Ca4CrKVkFEnVRGaHx", {"color": MUTED})]], size=13)
footer(s, 5)

# ---------- 6. CRE ----------
s = slide(
    "Best Workflow with CRE. One CRE workflow serves every note. Each cron run reads the engine's notes from "
    "Solana at finalized state, reads the Chainlink feed of every underlying at the finalized Ethereum block "
    "with a staleness check, and writes DON-signed reports through the keystone forwarder. One report runs an "
    "observation and everything after it. CRE is also the paying agent on Ethereum: a payout report has the "
    "engine lock exactly what a Sepolia holder is owed. The chart shows the devnet reports from the evidence, "
    "against the 300k compute cap at the top of the axis."
)
title(s, "Chainlink CRE is the calculation and paying agent", "Best Workflow with CRE")
roles = [
    ("Reads Solana", "Every DON node reads finalized engine state; the DON agrees on what is due"),
    ("Reads prices", "Chainlink feed of every underlying at the finalized Ethereum block, staleness-checked"),
    ("Settles", "One DON-signed report through the forwarder runs an observation and everything after it"),
    ("Pays Ethereum", "Payout reports lock each Sepolia holder's coupon, redemption or token for CCIP"),
]
y = Inches(1.95)
for head, sub in roles:
    dot(s, M, y + Inches(0.09), TEAL)
    text(s, M + Inches(0.35), y, Inches(5.0), Inches(1.0),
         [[(head, {"bold": True, "size": 17})], [(sub, {"size": 13, "color": SOFT})]], space_after=0)
    y += Inches(1.08)
cd = CategoryChartData()
cd.categories = ["Strike", "Obs 1", "Obs 2", "3 strikes", "Obs 1", "Obs 2"]
cd.add_series("Compute units", (51394, 68976, 87593, 63723, 84063, 99231))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.4), Inches(1.85), Inches(6.33), Inches(4.3), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Compute units per CRE report (cap 300,000)"
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size = Pt(13)
tp.runs[0].font.bold = True
tp.runs[0].font.color.rgb = TEXT
tp.runs[0].font.name = FONT
va = ch.value_axis
va.maximum_scale = 300000
va.minimum_scale = 0
va.major_unit = 100000
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = LINE
va.format.line.fill.background()
va.tick_labels.font.size = Pt(11)
va.tick_labels.font.color.rgb = MUTED
va.tick_labels.number_format = '#,##0'
va.tick_labels.number_format_is_linked = False
ca = ch.category_axis
ca.tick_labels.font.size = Pt(11)
ca.tick_labels.font.color.rgb = MUTED
ca.format.line.color.rgb = LINE
pl = ch.plots[0]
pl.gap_width = 70
pl.has_data_labels = True
dl = pl.data_labels
dl.number_format = '#,##0'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(11)
dl.font.color.rgb = TEXT
ser = pl.series[0]
for i, pt_col in enumerate([EMERALD] * 3 + [TEAL] * 3):
    p = ser.points[i]
    p.format.fill.solid()
    p.format.fill.fore_color.rgb = pt_col
text(s, Inches(6.4), Inches(6.2), Inches(6.33), Inches(0.3),
     [[("Phoenix on ETH", {"color": EMERALD, "bold": True}), ("   ·   ", {}),
       ("Worst-of ETH, BTC, SOL", {"color": TEAL, "bold": True})]], size=12, color=MUTED, align=PP_ALIGN.CENTER)
footer(s, 6)

# ---------- 7. Cross-chain ----------
s = slide(
    "Investors on Ethereum, over CCIP: one Ethereum address across three notes. On the FCN it subscribed 300 "
    "units from Sepolia; CRE autocalled and 306 tUSD went back. On the phoenix, CRE locked three coupons and the "
    "redemption in four payout reports, and 110 tUSD, the reference payoff, arrived with nobody acting. On the "
    "physically settled reverse convertible, the note knocked in and delivered 0.03824647 tETH, 100 units divided "
    "by the 2,614.62 strike, plus 5 tUSD of coupons. Outbound takes about a minute; inbound takes 35 to 40 "
    "minutes because of Ethereum finality."
)
title(s, "Subscribe and get paid on Solana or Ethereum", "Chainlink CCIP")
trio = [
    ("306", "tUSD", "Subscribe in", "300 units bought from Sepolia; CRE autocalled; paid back over CCIP", "FCN on ETH · DFaUHwaG…"),
    ("110", "tUSD", "Automatic payouts", "Three coupons and the redemption, locked by CRE, sent by the relay", "Phoenix · 8DtC17Ar…"),
    ("0.0382", "tETH", "Physical delivery", "0.03824647 tETH (100 ÷ 2,614.62 strike) plus 5 tUSD of coupons", "Reverse convertible · CviXTYq7…"),
]
tw = int((W - 2 * M - Inches(0.6)) / 3)
for i, (big, unit, head, sub, ref) in enumerate(trio):
    x = M + i * (tw + Inches(0.3))
    card(s, x, Inches(1.95), tw, Inches(3.15))
    text(s, x + Inches(0.3), Inches(2.15), tw - Inches(0.6), Inches(0.95),
         [[(big, {"size": 46, "bold": True, "color": VIOLET}), ("  " + unit, {"size": 20, "color": SOFT})]])
    text(s, x + Inches(0.3), Inches(3.15), tw - Inches(0.6), Inches(0.4), head, size=19, bold=True)
    text(s, x + Inches(0.3), Inches(3.6), tw - Inches(0.6), Inches(0.9), sub, size=14, color=SOFT)
    text(s, x + Inches(0.3), Inches(4.6), tw - Inches(0.6), Inches(0.35), ref, size=11, color=MUTED)
stats = [("~1 min", "Solana → Ethereum"), ("35–40 min", "Ethereum → Solana"), ("0", "manual steps for payouts")]
for i, (big, lab) in enumerate(stats):
    x = M + i * (tw + Inches(0.3))
    text(s, x, Inches(5.4), tw, Inches(0.6), big, size=30, bold=True, color=TEXT)
    text(s, x, Inches(6.0), tw, Inches(0.35), lab, size=14, color=MUTED)
footer(s, 7)

# ---------- 8. Evidence ----------
s = slide(
    "Two whole lives on Solana devnet, run by CRE with nobody acting except one investor selling part of a "
    "position. The ETH phoenix paid 525, 420 and 105 USDC for 500, 400 and 100 units. The worst-of phoenix on "
    "ETH, BTC and SOL computed the worst performance on-chain from three Chainlink prices per report and paid "
    "636 and 424 USDC for 600 and 400 units. Both match the reference payoff exactly."
)
title(s, "Two full lifecycles, exactly the reference payoff", "Evidence")
cw2 = int((W - 2 * M - Inches(0.4)) / 2)
for i, (head, sub, rows) in enumerate([
    ("Phoenix on ETH", "8SvcVvQ8… · paid 525, 420, 105 USDC", [
        ["Event", "Transaction", "Compute units"],
        ["CRE fixes the strike", "2VxTRWLp…", "51,394"],
        ["Investor sells 100 of 600 units", "3umFWgcG…", "34,775"],
        ["Observation 1: coupon", "Ju2Wo4k5…", "68,976"],
        ["Observation 2: autocall, redemption", "4hySSWXN…", "87,593"],
    ]),
    ("Worst-of ETH, BTC, SOL", "3o6EuEWk… · paid 636, 424 USDC", [
        ["Event", "Transaction", "Compute units"],
        ["CRE fixes three strikes", "2JoaJh49…", "63,723"],
        ["Observation 1: worst 100%, coupon", "3GL2E2Q7…", "84,063"],
        ["Observation 2: autocall, redemption", "42Q2F5oR…", "99,231"],
    ]),
]):
    x = M + i * (cw2 + Inches(0.4))
    text(s, x, Inches(1.95), cw2, Inches(0.4), head, size=20, bold=True)
    text(s, x, Inches(2.4), cw2, Inches(0.3), sub, size=13, color=MUTED)
    table(s, x, Inches(2.9), cw2, rows, [Inches(3.3), Inches(1.2), cw2 - Inches(4.5)], size=13)
text(s, M, Inches(5.65), W - 2 * M, Inches(0.8),
     "Worst-of: each CRE report carried three Chainlink prices; the engine computed perfK = min(obsK_i / initialLevel_i) "
     "on-chain and tested every barrier against it. Full transaction ids: docs/SUBMISSION.md and docs/crosschain.md.",
     size=13, color=SOFT)
footer(s, 8)

# ---------- 9. AI + templates ----------
s = slide(
    "Design with AI: a full phoenix term sheet becomes the right parameters in about seven seconds. Ten percent a "
    "year paid quarterly becomes 2.5 percent a period; the worst of ETH and BTC becomes a two-asset basket. A "
    "request missing terms gets a question back instead of a guess, and asking for a paying-agent fee comes back "
    "as a validated workflow change. The template library has six templates on devnet, each signed by its author "
    "and recompiled against its definition address; forged or mismatched entries are rejected."
)
title(s, "Design by sentence, issue from a signed library", "AI and templates")
lx, lw = M, Inches(6.0)
card(s, lx, Inches(1.95), lw, Inches(1.3), fill=BG2, line=TEAL)
text(s, lx + Inches(0.3), Inches(2.05), lw - Inches(0.6), Inches(1.1),
     [[("PROMPT", {"size": 11, "bold": True, "color": TEAL})],
      [("“12 month phoenix on the worst of ETH and BTC, 10% p.a. paid quarterly, 70% coupon barrier with memory…”",
        {"size": 14, "italic": True})]], space_after=2, anchor=MSO_ANCHOR.MIDDLE)
ai = [
    ("~7 s", "term sheet to parameters"),
    ("10% p.a. → 2.5%", "per quarterly period"),
    ("2 underlyings", "become a worst-of basket"),
    ("Asks", "when terms are missing, never invents"),
]
for i, (big, lab) in enumerate(ai):
    x = lx + (i % 2) * Inches(3.1)
    y = Inches(3.55) + (i // 2) * Inches(1.25)
    text(s, x, y, Inches(2.9), Inches(0.5), big, size=22, bold=True, color=TEAL)
    text(s, x, y + Inches(0.5), Inches(2.9), Inches(0.5), lab, size=13, color=SOFT)
rx, rw = Inches(7.1), W - M - Inches(7.1)
card(s, rx, Inches(1.95), rw, Inches(4.45))
text(s, rx + Inches(0.3), Inches(2.15), rw - Inches(0.6), Inches(0.4), "Template library: six on devnet", size=18, bold=True)
tpls = ["Worst-of phoenix (ETH, BTC, SOL)", "Worst-of FCN (ETH, BTC)", "BTC snowball", "90%-protected SOL note",
        "Step-down phoenix (custom workflow)", "BTC phoenix with a servicing fee (custom)"]
y = Inches(2.7)
for t in tpls:
    dot(s, rx + Inches(0.3), y + Inches(0.08), EMERALD, d=Inches(0.12))
    text(s, rx + Inches(0.55), y, rw - Inches(0.85), Inches(0.35), t, size=14)
    y += Inches(0.43)
text(s, rx + Inches(0.3), Inches(5.4), rw - Inches(0.6), Inches(0.8),
     "Each entry is signed by its author and recompiled against its definition address; forgeries are rejected.",
     size=12, color=MUTED)
footer(s, 9)

# ---------- 10. Tests ----------
s = slide(
    "The products end-to-end suite runs twelve lifecycles on a local validator: all five products, worst-of "
    "baskets, physical delivery and the two custom templates, with a partial transfer mid-life. Every payout "
    "equals the expected payoff. The cross-chain suite uses a mock CCIP router and offramp with the real PDAs: "
    "a Sepolia subscription, a late one refunded, and the autocall paid through the CRE path: payout report, "
    "outbox, flush_outbox, ccip_send."
)
title(s, "Every payout checked against a reference payoff", "Tests")
tests = [
    ("12", "product lifecycles", "All five products, worst-of, physical delivery, two custom templates; every payout equals the reference payoff"),
    ("3", "cross-chain scenarios", "Subscribe via ccip_receive; late subscription refunded; autocall paid through payout report → outbox → flush_outbox"),
    ("< 300k", "compute per CRE report", "Every report in the e2e suite fits CRE's cap (max 98k there)"),
    ("6", "signed templates", "Library rejects forged signatures and mismatched workflows, checked against production"),
]
tw4 = int((W - 2 * M - Inches(0.3)) / 2)
for i, (big, lab, sub) in enumerate(tests):
    x = M + (i % 2) * (tw4 + Inches(0.3))
    y = Inches(1.95) + (i // 2) * Inches(2.05)
    card(s, x, y, tw4, Inches(1.8))
    text(s, x + Inches(0.3), y + Inches(0.2), Inches(1.9), Inches(0.8), big, size=40, bold=True, color=EMERALD)
    text(s, x + Inches(2.25), y + Inches(0.3), tw4 - Inches(2.5), Inches(0.5), lab, size=18, bold=True)
    text(s, x + Inches(0.3), y + Inches(1.0), tw4 - Inches(0.6), Inches(0.75), sub, size=13, color=SOFT)
text(s, M, Inches(6.0), W - 2 * M, Inches(0.65),
     "Also: unit tests for every product's BPMN round-trip, min / max, worst-of payoff and the AI pipeline; Foundry "
     "tests for the Sepolia faucet; the CRE workflow compiles with cre-compile.", size=13, color=MUTED)
footer(s, 10)

# ---------- 11. Limits and next ----------
s = slide(
    "Be direct about the limits. Inbound CCIP is slow, so a book that takes Ethereum subscriptions stays open at "
    "least 40 minutes. The DON executor runs out of heap after our receiver succeeds, so a relayer re-executes "
    "those messages manually. CRE locks payouts but cannot submit the CCIP send itself because the router needs "
    "address lookup tables, so a permissionless relay delivers them. The tokens are our own CCIP burn-mint test "
    "tokens: we control the lane and the faucet, and a note points at mint addresses, so real USDC needs no "
    "engine change. Next: deploy to the DON, note units as SPL tokens, more chains."
)
title(s, "What is not done yet, and what comes next", "Limits and next")
lw2 = int((W - 2 * M - Inches(0.4)) / 2)
limits = [
    ("Inbound CCIP is slow", "Ethereum → Solana takes 35–40 min; Ethereum-open books stay open at least 40 min"),
    ("DON executor and our receiver", "The offramp runs out of heap after our receiver succeeds; our relayer re-executes"),
    ("CRE decides, a relay sends", "CRE writes can't use lookup tables; the relay cannot change amount or destination"),
    ("Test tokens", "Our own CCIP tokens: we control the lane and the faucet. Notes point at mint addresses, so USDC needs no engine change"),
    ("CRE simulator", "Runs with --broadcast until DON deploy access; the engine accepts the staging forwarder"),
]
nexts = [
    "cre workflow deploy to the DON; log triggers on engine events",
    "Note units as SPL tokens (Token-2022), transferable over CCIP",
    "More CCIP chains (Base, Arbitrum); a smaller receiver footprint",
    "Issuer reserves from Ethereum in the app",
    "Investor allowlists, template ratings, issuer profiles",
    "Flow's Canton target for private institutional issuance",
]
text(s, M, Inches(1.9), lw2, Inches(0.4), "Limits", size=18, bold=True, color=VIOLET)
y = Inches(2.35)
for head, sub in limits:
    text(s, M, y, lw2, Inches(0.8), [[(head, {"bold": True, "size": 15})], [(sub, {"size": 12, "color": SOFT})]], space_after=0)
    y += Inches(0.82)
rx2 = M + lw2 + Inches(0.4)
card(s, rx2, Inches(1.9), lw2, Inches(4.6))
text(s, rx2 + Inches(0.3), Inches(2.1), lw2 - Inches(0.6), Inches(0.4), "Next", size=18, bold=True, color=EMERALD)
y = Inches(2.65)
for t in nexts:
    dot(s, rx2 + Inches(0.3), y + Inches(0.08), EMERALD, d=Inches(0.12))
    text(s, rx2 + Inches(0.55), y, lw2 - Inches(0.85), Inches(0.55), t, size=14)
    y += Inches(0.6)
footer(s, 11)

# ---------- 12. Close ----------
s = slide(
    "One Solana program runs every note as a BPMN workflow; one CRE workflow observes them all and pays holders "
    "on Solana and Ethereum; CCIP moves the money; anyone can issue from the library. The links are in the submission."
)
text(s, M, Inches(1.2), Inches(11.5), Inches(2.4), [
    [("One Solana program runs every note.", {"size": 36, "bold": True})],
    [("One CRE workflow observes and pays them all.", {"size": 36, "bold": True, "color": TEAL})],
    [("CCIP moves the money.", {"size": 36, "bold": True, "color": VIOLET})],
], space_after=4)
links = [
    ("Site", "sp.stratoslab.app"),
    ("App", "sp.stratoslab.app/app"),
    ("Code", "github.com/adappterxyz/StratosNotes"),
    ("Docs", "docs/SUBMISSION.md · docs/DEMO.md · docs/crosschain.md"),
]
y = Inches(4.3)
for lab, val in links:
    text(s, M, y, Inches(1.2), Inches(0.4), lab.upper(), size=13, bold=True, color=EMERALD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + Inches(1.3), y, Inches(9), Inches(0.4), val, size=18, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.52)
text(s, M, Inches(6.6), Inches(8), Inches(0.35), "StratosNotes · TOKEN2049 Origins", size=12, color=MUTED)

prs.save(OUT)
print(OUT)
